"""A content brief is read into one request and run through the existing flow."""

import base64
import io
import json
from types import SimpleNamespace
from unittest.mock import patch

import pytest

from app.services import brief as brief_service
from app.services.brief import BriefError, curate_brief, read_brief_text
from tests.plans import discovery_plan, success_result


def _b64(data: bytes) -> str:
    return base64.b64encode(data).decode()


def _pdf_with_text(text: str) -> bytes:
    """A one-page PDF whose page draws `text` — enough for pypdf to read."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return out.getvalue()


def _docx_with(paragraph: str, cells) -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph(paragraph)
    table = document.add_table(rows=1, cols=len(cells))
    for cell, value in zip(table.rows[0].cells, cells):
        cell.text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


# ── reading the file ────────────────────────────────────────────────


def test_a_text_brief_is_read_as_it_is():
    text = read_brief_text("brief.md", _b64(b"# Campaign\n\nFashion creators in Riyadh\n"))
    assert text == "# Campaign\nFashion creators in Riyadh"


def test_a_pdf_brief_is_read():
    assert "Fashion creators in Riyadh" in read_brief_text(
        "Brief.PDF", _b64(_pdf_with_text("Fashion creators in Riyadh"))
    )


def test_a_word_brief_keeps_its_tables():
    """Market, platform and size often live in a table, not the prose."""
    text = read_brief_text(
        "brief.docx", _b64(_docx_with("Launch of the abaya line", ["Market", "Saudi Arabia"]))
    )
    assert "Launch of the abaya line" in text
    assert "Market | Saudi Arabia" in text


@pytest.mark.parametrize("filename, content, code", [
    ("brief.pptx", _b64(b"x"), "brief_unsupported_type"),
    ("brief", _b64(b"x"), "brief_unsupported_type"),
    ("brief.txt", "not base64!!", "brief_unreadable"),
    ("brief.txt", _b64(b"   \n  "), "brief_no_text"),
    ("brief.pdf", _b64(b"%PDF-1.4 broken"), "brief_unreadable"),
])
def test_a_brief_that_cannot_be_read_says_why(filename, content, code):
    with pytest.raises(BriefError) as err:
        read_brief_text(filename, content)
    assert err.value.code == code


def test_a_scanned_pdf_has_no_text_to_read():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    out = io.BytesIO()
    writer.write(out)
    with pytest.raises(BriefError) as err:
        read_brief_text("scan.pdf", _b64(out.getvalue()))
    assert err.value.code == "brief_no_text"


def test_an_oversized_brief_is_refused(monkeypatch):
    monkeypatch.setattr(brief_service, "MAX_BRIEF_BYTES", 10)
    with pytest.raises(BriefError) as err:
        read_brief_text("brief.txt", _b64(b"x" * 11))
    assert err.value.code == "brief_too_large"


# ── reading it into research tracks ─────────────────────────────────


def _model_says(content: str):
    return SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(
        create=lambda **kw: SimpleNamespace(
            choices=[SimpleNamespace(message=SimpleNamespace(content=content))]
        )
    )))


TWO_TRACKS = json.dumps({
    "summary": "Recruitment narratives reaching two groups.",
    "tracks": [
        {"group": "Group A", "question": "creators in Country X on TikTok posting about work offers",
         "countries": ["Country X", "Country X"], "languages": ["Language 1"],
         "platforms": ["TikTok"], "inclusion": ["shows location", "  "]},
        {"group": "Group B", "question": "creators in Country Y on Instagram",
         "countries": ["Country Y"], "languages": [], "platforms": ["Instagram"], "inclusion": []},
        {"group": "No question", "question": "  "},
        "not a track",
    ],
    "gaps": ["No language given for Group B"],
})


def test_each_group_in_the_brief_is_its_own_track():
    with patch("app.services.brief.OpenAI", return_value=_model_says(TWO_TRACKS)):
        brief = curate_brief("the brief", "", filename="b.pdf", openai_key="sk", model="m")
    assert brief.filename == "b.pdf"
    assert [t.group for t in brief.tracks] == ["Group A", "Group B"]
    first = brief.tracks[0]
    assert first.question == "creators in Country X on TikTok posting about work offers"
    assert first.countries == ["Country X"]          # de-duplicated
    assert first.inclusion == ["shows location"]     # blanks dropped
    assert brief.gaps == ["No language given for Group B"]
    assert brief.summary.startswith("Recruitment narratives")


def test_a_track_without_a_group_name_is_still_kept():
    said = json.dumps({"tracks": [{"question": "creators in Country X"}]})
    with patch("app.services.brief.OpenAI", return_value=_model_says(said)):
        brief = curate_brief("b", "", filename="b.pdf", openai_key="sk", model="m")
    assert [t.group for t in brief.tracks] == ["Track 1"]


def test_the_operators_note_reaches_the_model():
    sent = {}

    def create(**kw):
        sent.setdefault("messages", kw["messages"])  # the brief reading is the first call
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=TWO_TRACKS))])

    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    with patch("app.services.brief.OpenAI", return_value=client):
        curate_brief("the brief", "TikTok only", filename="b.pdf", openai_key="sk", model="m")
    user = sent["messages"][1]["content"]
    assert "TikTok only" in user and "the brief" in user


@pytest.mark.parametrize("target, code", [
    ({"side_effect": RuntimeError("down")}, "brief_failed"),
    ({"return_value": _model_says('{"tracks": []}')}, "brief_no_tracks"),
    ({"return_value": _model_says("not json")}, "brief_failed"),
])
def test_a_brief_that_cannot_be_read_into_tracks_says_so(target, code):
    """Nothing is guessed past: the operator approves tracks, so no tracks is an error."""
    with patch("app.services.brief.OpenAI", **target):
        with pytest.raises(BriefError) as err:
            curate_brief("b", "", filename="b.pdf", openai_key="sk", model="m")
    assert err.value.code == code


def test_nothing_about_any_brief_is_written_into_the_prompt():
    """Groups, places and platforms come from the brief, never from the code."""
    prompt = brief_service.BRIEF_SYSTEM
    assert "What the brief does not give is not mentioned at all" in prompt
    assert "Never invent" in prompt
    assert "never merge two groups" in prompt
    assert "overrides the brief" in prompt


# ── the endpoints ───────────────────────────────────────────────────


BRIEF = {"filename": "brief.md", "content_base64": _b64(b"Research two groups.")}


def _tracks():
    return brief_service.parse_research_brief(json.loads(TWO_TRACKS), filename="brief.md")


@patch("app.api.v1.endpoints.research.curate_upload")
@patch("app.api.v1.endpoints.research.generate_research_plan")
def test_a_brief_returns_tracks_for_approval_and_searches_nothing(mock_gen, mock_curate, client):
    mock_curate.return_value = _tracks()

    response = client.post("/ask", json={"prompt": "focus on TikTok", "brief": BRIEF})

    assert response.status_code == 200
    body = response.json()
    assert [t["group"] for t in body["research_brief"]["tracks"]] == ["Group A", "Group B"]
    assert body["awaiting_approval"] is True
    mock_gen.assert_not_called()
    filename, text, note = mock_curate.call_args.args
    assert (filename, text, note) == ("brief.md", "Research two groups.", "focus on TikTok")

    from app.services import conversations as conversation_service
    turns = conversation_service.store.history(body["conversation_id"], "local-dev")
    assert turns[0].content == "Uploaded the brief \u201cbrief.md\u201d \u2014 focus on TikTok"
    assert json.loads(turns[1].content)["research_brief"]["tracks"][1]["group"] == "Group B"


@patch("app.api.v1.endpoints.research.curate_upload")
def test_a_brief_needs_no_note(mock_curate, client):
    mock_curate.return_value = _tracks()
    response = client.post("/ask", json={"brief": BRIEF})
    assert response.status_code == 200
    assert mock_curate.call_args.args[2] is None


@patch("app.api.v1.endpoints.research.generate_research_plan")
def test_an_approved_track_runs_as_an_ordinary_question(mock_gen, client):
    mock_gen.return_value = success_result(discovery_plan(), "discovery")

    response = client.post("/ask", json={"prompt": "creators in Country Y on Instagram", "track": "Group B"})

    assert response.status_code == 200
    assert mock_gen.call_args.args[0] == "creators in Country Y on Instagram"
    from app.services import conversations as conversation_service
    turns = conversation_service.store.history(response.json()["conversation_id"], "local-dev")
    assert turns[0].content == "Track \u201cGroup B\u201d: creators in Country Y on Instagram"


@patch("app.api.v1.endpoints.research.generate_research_plan")
def test_a_typed_question_is_untouched(mock_gen, client):
    mock_gen.return_value = success_result(discovery_plan(), "discovery")

    with patch("app.api.v1.endpoints.research.curate_upload") as curate:
        response = client.post("/ask", json={"prompt": "Find modest fashion creators in Saudi Arabia"})

    assert response.status_code == 200
    assert "research_brief" not in response.json()
    assert mock_gen.call_args.args[0] == "Find modest fashion creators in Saudi Arabia"
    curate.assert_not_called()
    from app.services import conversations as conversation_service
    turns = conversation_service.store.history(response.json()["conversation_id"], "local-dev")
    assert turns[0].content == "Find modest fashion creators in Saudi Arabia"


def test_an_empty_prompt_without_a_brief_is_still_refused(client):
    assert client.post("/ask", json={"prompt": ""}).status_code == 422
    assert client.post("/ask", json={}).status_code == 422


@patch("app.api.v1.endpoints.research.generate_research_plan")
def test_an_unreadable_brief_is_400_before_anything_runs(mock_gen, client):
    response = client.post("/ask", json={
        "prompt": "creators for this",
        "brief": {"filename": "brief.pptx", "content_base64": _b64(b"x")},
    })
    assert response.status_code == 400
    assert response.json()["code"] == "brief_unsupported_type"
    mock_gen.assert_not_called()


@patch("app.api.v1.endpoints.research.curate_upload",
       side_effect=BriefError("No research tracks could be read from the brief.", "brief_no_tracks"))
def test_a_brief_with_no_tracks_is_said_plainly(mock_curate, client):
    response = client.post("/ask", json={"brief": BRIEF})
    assert response.status_code == 422
    assert response.json()["code"] == "brief_no_tracks"


@patch("app.api.v1.endpoints.research.curate_upload")
@patch("app.api.v1.endpoints.research.generate_research_plan")
def test_the_stream_reads_the_brief_and_stops_for_approval(mock_gen, mock_curate, client):
    mock_curate.return_value = _tracks()

    with client.stream("POST", "/ask/stream", json={"brief": BRIEF}) as response:
        assert response.status_code == 200
        lines = list(response.iter_lines())

    events = [l[7:] for l in lines if l.startswith("event: ")]
    data = [json.loads(l[6:]) for l in lines if l.startswith("data: ")]
    assert events == ["progress", "progress", "done"]
    assert data[1]["stage"] == "brief"
    assert [t["group"] for t in data[-1]["research_brief"]["tracks"]] == ["Group A", "Group B"]
    mock_gen.assert_not_called()


@patch("app.api.v1.endpoints.research.curate_upload",
       side_effect=BriefError("The brief could not be read into research tracks. Try again.", "brief_failed"))
def test_the_stream_says_when_the_brief_could_not_be_read(mock_curate, client):
    with client.stream("POST", "/ask/stream", json={"brief": BRIEF}) as response:
        lines = list(response.iter_lines())
    events = [l[7:] for l in lines if l.startswith("event: ")]
    assert events[-1] == "failed"
    assert "brief_failed" in lines[-1] or any("brief_failed" in l for l in lines)
