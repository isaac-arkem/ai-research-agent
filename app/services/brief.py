"""A content brief, read into the request the agent already understands.

An operator uploads a brand's brief — a PDF or Word file of objectives,
audience, markets, deliverables — and wants creators for it. The agent has
no separate brief pipeline and should not grow one: it already turns
"<kind of creator> in <country> on <platform>, under <size>" into a
search, filters it, and drops brands. So the brief is read into
research TRACKS — one per group the brief asks about, each with the request the
existing flow will run — and the operator approves them before anything is
searched. An approved track runs exactly as if the operator had typed its
question. A brief that names creators to be like becomes "similar to @x", and
the existing seed flow looks them up first and searches from what it finds.

Two steps, both here:
  read_brief_text  — the file's words. No model.
  curate_brief     — the words, as research tracks and gaps. The model decides
                     what the brief asks for; nothing here matches on keywords,
                     and no group, place or platform is written into the code.
"""

from __future__ import annotations

import base64
import binascii
import io
import logging
from typing import List, Optional

from openai import OpenAI

from app.models.domain import ResearchBrief, ResearchTrack

from app.utils.guards import MAX_PROMPT_LENGTH
from app.utils.json_extract import extract_json_object

logger = logging.getLogger(__name__)

# A brief is a few pages. 3 MB of file covers a designed PDF with images;
# the words are what matter, and 200k characters is far more than any brief.
MAX_BRIEF_BYTES = 3_000_000
MAX_BRIEF_CHARS = 200_000

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md")


class BriefError(ValueError):
    """A brief that cannot be read. `code` goes back to the client."""

    def __init__(self, message: str, code: str):
        super().__init__(message)
        self.code = code


def _extension(filename: str) -> str:
    name = (filename or "").strip().lower()
    return name[name.rfind("."):] if "." in name else ""


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _docx_text(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    # Briefs put the useful part in tables as often as in prose: market,
    # platform, follower range, audience.
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def read_brief_text(filename: str, content_base64: str) -> str:
    """The words in an uploaded brief, or BriefError saying why not."""
    ext = _extension(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise BriefError(
            "Upload the brief as a PDF, Word (.docx) or text file.",
            "brief_unsupported_type",
        )
    try:
        data = base64.b64decode(content_base64 or "", validate=True)
    except (binascii.Error, ValueError):
        raise BriefError("The brief file could not be read.", "brief_unreadable")
    if not data:
        raise BriefError("The brief file is empty.", "brief_empty")
    if len(data) > MAX_BRIEF_BYTES:
        raise BriefError("The brief file is larger than 3 MB.", "brief_too_large")

    try:
        if ext == ".pdf":
            text = _pdf_text(data)
        elif ext == ".docx":
            text = _docx_text(data)
        else:
            text = data.decode("utf-8", errors="replace")
    except Exception as exc:  # a corrupt file is the operator's to fix, not a crash
        logger.warning("brief: could not read %s: %s", ext, exc)
        raise BriefError("The brief file could not be read.", "brief_unreadable")

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not text:
        # A scanned PDF is pictures of words; there is nothing to read.
        raise BriefError(
            "No text found in the brief. A scanned PDF has none — upload the original document.",
            "brief_no_text",
        )
    return text[:MAX_BRIEF_CHARS]


BRIEF_SYSTEM = """You read a content brief and turn it into a research brief for a creator-research agent. The operator approves it before anything is searched.

The agent finds social media creators and their content. It runs one request at a time, written the way an operator types it: a plain sentence naming the kind of creator, and then only what the brief gives — the country, the platform, the size, who the creators are, what they post about. A reference creator is written "similar to @handle". What the brief does not give is not mentioned at all: never write "any platform", "any size", "any country" or the like, not even once — the agent asks for what it needs.
It searches, scrapes and filters on its own. Your job is only the research brief.

TRACKS. When the brief asks for research on different groups of people, audiences or situations, each one is its own track — never merge two groups into one request, because their results must stay apart. When it asks for one, there is one track. Name each group in the brief's own words.

For each track, in this order:
- group: who this track is about, in the brief's words;
- countries, languages, platforms: exactly as the brief gives them for this track, empty when it does not say;
- inclusion: everything concrete the brief says marks the people or posts of this track — who they are, their situation, what they are offered or promised and through which route, what the brief says to record — short, one per line, in the brief's words. A brief describing personas or scenarios still describes the real people and posts to look for: use those descriptions;
- question: the request the agent will run, one or two plain sentences. The question is ALL the agent sees, so it must carry everything above that a search can use: this track's countries, languages, platforms, size, and every inclusion criterion and exclusion. It asks for what to find, never for what to record.
  Write "similar to @name" ONLY when the brief names a specific real account as an example to be like — then the agent looks that account up and searches from it. Otherwise never write "similar to".

Leave out what is not a search: budget, fees, deliverables, schedules, team roles, campaign mechanics, the client's history.

GAPS. List what the research needs and the brief does not say — a group with no country (or only a region, when a search needs countries), no language, no platform, no criteria for who counts. Short, one per line. Empty when nothing is missing.

SUMMARY. One sentence: what the brief asks the research to find.

Never invent. What the brief does not say stays empty and goes in the gaps.

The operator may add a note. It overrides the brief where they differ.

Write in English.

Return JSON only:
{"summary": "...", "tracks": [{"group": "...", "countries": [], "languages": [], "platforms": [], "inclusion": [], "question": "..."}], "gaps": []}"""

# Bounds on what the model returns, not on what a brief may contain.
_MAX_TRACKS = 12
_MAX_ITEMS = 25


def _texts(value) -> List[str]:
    out: List[str] = []
    for raw in value if isinstance(value, list) else []:
        text = " ".join(str(raw or "").split())
        if text and text not in out:
            out.append(text[:300])
    return out[:_MAX_ITEMS]


def parse_research_brief(parsed: dict, *, filename: str) -> ResearchBrief:
    """The model's JSON as a ResearchBrief, keeping only well-formed tracks."""
    parsed = parsed if isinstance(parsed, dict) else {}
    tracks: List[ResearchTrack] = []
    for row in parsed.get("tracks") if isinstance(parsed.get("tracks"), list) else []:
        if not isinstance(row, dict):
            continue
        question = " ".join(str(row.get("question") or "").split())[:MAX_PROMPT_LENGTH]
        if not question:
            continue
        group = " ".join(str(row.get("group") or "").split())[:200]
        tracks.append(ResearchTrack(
            group=group or f"Track {len(tracks) + 1}",
            question=question,
            countries=_texts(row.get("countries")),
            languages=_texts(row.get("languages")),
            platforms=_texts(row.get("platforms")),
            inclusion=_texts(row.get("inclusion")),
        ))
    return ResearchBrief(
        filename=filename,
        summary=" ".join(str(parsed.get("summary") or "").split())[:600],
        tracks=tracks[:_MAX_TRACKS],
        gaps=_texts(parsed.get("gaps")),
    )


def curate_brief(
    text: str,
    note: str,
    *,
    filename: str,
    openai_key: str,
    model: str,
    timeout: float = 60.0,
) -> ResearchBrief:
    """The brief as research tracks, or BriefError when it could not be read.

    No fallback to running the raw brief: the point of this step is a brief
    the operator approves, and a failure is better said than guessed past.
    """
    note = (note or "").strip()
    try:
        client = OpenAI(api_key=openai_key, timeout=timeout)
        response = client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system", "content": BRIEF_SYSTEM},
                {"role": "user", "content": (
                    (f"The operator's note: {note}\n\n" if note else "")
                    + f"The brief ({filename}):\n{text}"
                )},
            ],
            temperature=0,
            response_format={"type": "json_object"},
        )
        parsed = extract_json_object(response.choices[0].message.content or "") or {}
    except Exception as exc:
        logger.warning("brief: reading %s failed: %s", filename, exc)
        raise BriefError("The brief could not be read into research tracks. Try again.", "brief_failed")

    brief = parse_research_brief(parsed, filename=filename)
    if not brief.tracks:
        logger.warning("brief: %s produced no tracks", filename)
        raise BriefError("No research tracks could be read from the brief.", "brief_no_tracks")
    logger.info("brief: %s -> %d track(s): %s", filename, len(brief.tracks),
                [t.group for t in brief.tracks])
    return brief


def brief_message(filename: str, note: Optional[str] = None) -> str:
    """What the thread stores as the operator's turn when they upload a brief."""
    note = (note or "").strip()
    return f"Uploaded the brief \u201c{filename}\u201d" + (f" \u2014 {note}" if note else "")


def track_message(group: str, question: str) -> str:
    """What the thread stores when the operator runs an approved track, so the
    results that follow stay labelled with their group."""
    return f"Track \u201c{group}\u201d: {question}"


def curate_upload(
    filename: str,
    text: str,
    note: Optional[str],
    *,
    settings,
) -> ResearchBrief:
    """curate_brief with the agent's settings."""
    return curate_brief(
        text, note or "",
        filename=filename,
        openai_key=settings.openai_api_key,
        model=getattr(settings, "grounding_escalation_model", None) or settings.research_agent_model,
        timeout=max(float(getattr(settings, "search_timeout", 15.0)), 60.0),
    )
